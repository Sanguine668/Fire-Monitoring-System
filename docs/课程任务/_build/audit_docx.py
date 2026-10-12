"""对生成的 Word 文档做结构化校验（标题层级、表格、图片占位）。"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[3]
FILES = [
    ROOT / "docs/课程任务/实验5-软件项目进度管理/实验5_软件项目进度管理.docx",
    ROOT / "docs/课程任务/实验6-软件项目成本管理/实验6_软件项目成本管理.docx",
]


def audit(path: Path) -> None:
    doc = Document(str(path))
    heads: dict[str, int] = {}
    placeholders = 0
    short_tables = 0
    cells = 0
    for para in doc.paragraphs:
        style = para.style.name
        if style.startswith("Heading"):
            heads[style] = heads.get(style, 0) + 1
        if "此处插入图片" in para.text:
            placeholders += 1
    for table in doc.tables:
        cells += len(table.rows) * len(table.columns)
        if len(table.rows) == 1 and len(table.columns) == 1:
            text = table.rows[0].cells[0].text
            if "此处插入图片" in text:
                placeholders += 1
            elif not text.strip():
                short_tables += 1
    print(f"--- {path.name} ---")
    print(f"  大小: {path.stat().st_size / 1024:.1f} KB")
    print(f"  段落 {len(doc.paragraphs)}  表格 {len(doc.tables)}  单元格 {cells}")
    print(f"  标题层级: {heads}")
    print(f"  图片占位 {placeholders} 处；空占位框 {short_tables} 个")


def dump_outline(path: Path) -> None:
    doc = Document(str(path))
    print(f"--- {path.name} 标题结构 ---")
    for para in doc.paragraphs:
        style = para.style.name
        text = para.text.strip()
        if not text:
            continue
        if style == "Heading 1":
            print(f"  [H1] {text}")
        elif style == "Heading 2":
            print(f"    [H2] {text}")
        elif style == "Heading 3":
            print(f"      [H3] {text}")


def main() -> None:
    for path in FILES:
        if path.is_file():
            audit(path)
        else:
            print(f"MISS {path}")


if __name__ == "__main__":
    main()
