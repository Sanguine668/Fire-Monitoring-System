"""把实验 5、6 的 Markdown 转成 Word 文档，图片位置留成占位框。

用法：ai\\.venv\\Scripts\\python docs\\课程任务\\_build\\build_exp56_docx.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent))
from common import (  # noqa: E402
    ALT_FILL,
    BODY_SIZE,
    USABLE_WIDTH_CM,
    _set_cell_margins,
    _set_table_widths,
    add_page_number_footer,
    add_table,
    add_toc_field,
    caption,
    configure_styles,
    h1,
    h2,
    h3,
    set_run_font,
)
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

ROOT = Path(__file__).resolve().parents[3]
TASKS = [
    (
        ROOT / "docs/课程任务/实验5-软件项目进度管理/实验5_软件项目进度管理.md",
        ROOT / "docs/课程任务/实验5-软件项目进度管理/实验5_软件项目进度管理.docx",
    ),
    (
        ROOT / "docs/课程任务/实验6-软件项目成本管理/实验6_软件项目成本管理.md",
        ROOT / "docs/课程任务/实验6-软件项目成本管理/实验6_软件项目成本管理.docx",
    ),
]

PLACEHOLDER_HINT = "（截图位置）"


def make_box(doc, fill: str | None = None):
    """建一个与正文同宽的 1×1 框（提示框 / 代码框 / 图片占位框）。"""
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_widths(table, [USABLE_WIDTH_CM])
    _set_cell_margins(table)
    cell = table.rows[0].cells[0]
    cell.text = ""
    if fill:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)
    return table, cell


def add_inline(para, text: str, size=None) -> None:
    """把 **加粗** 与 `代码` 解析成 run。"""
    pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
    for piece in pattern.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            run = para.add_run(piece[2:-2])
            set_run_font(run, size=size or BODY_SIZE, bold=True)
        elif piece.startswith("`") and piece.endswith("`"):
            run = para.add_run(piece[1:-1])
            set_run_font(run, cjk="Consolas", latin="Consolas", size=Pt(9.5))
        else:
            run = para.add_run(piece)
            set_run_font(run, size=size or BODY_SIZE)


def body_paragraph(doc, text: str, indent: bool = True, size=None):
    para = doc.add_paragraph()
    if indent:
        para.paragraph_format.first_line_indent = Cm(0.74)
    add_inline(para, text, size=size)
    return para


def note_block(doc, text: str) -> None:
    """引用/提示块：用浅色底纹的单元格呈现。"""
    _table, cell = make_box(doc, ALT_FILL)
    para = cell.paragraphs[0]
    para.paragraph_format.first_line_indent = Cm(0)
    add_inline(para, text, size=Pt(9.5))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def image_placeholder(doc, text: str) -> None:
    """图片占位框：一个带边框的空框 + 图题。"""
    clean = text.strip().strip("*").strip()
    _table, cell = make_box(doc)
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = Cm(0)
    for _ in range(3):
        para.add_run("\n")
    run = para.add_run(f"【此处插入图片】{clean}")
    set_run_font(run, size=Pt(10.5), bold=True, color=RGBColor(0x80, 0x80, 0x80))
    for _ in range(3):
        para.add_run("\n")
    caption(doc, clean)


def code_block(doc, lines: list[str]) -> None:
    _table, cell = make_box(doc, "F7F7F7")
    para = cell.paragraphs[0]
    para.paragraph_format.first_line_indent = Cm(0)
    para.paragraph_format.line_spacing = 1.15
    for idx, line in enumerate(lines):
        run = para.add_run(line)
        set_run_font(run, cjk="Consolas", latin="Consolas", size=Pt(9.5))
        if idx != len(lines) - 1:
            para.add_run("\n")


def parse_table_rows(lines: list[str], start: int):
    rows = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        rows.append(lines[idx].strip())
        idx += 1
    return rows, idx


def split_row(row: str) -> list[str]:
    return [cell.strip() for cell in row.strip().strip("|").split("|")]


def table_widths(col_count: int) -> list[float]:
    if col_count == 2:
        return [4.4, 11.0]
    if col_count == 3:
        return [3.0, 6.2, 6.2]
    if col_count == 4:
        return [2.2, 5.6, 3.8, 3.8]
    if col_count == 5:
        return [1.8, 5.0, 3.0, 2.8, 2.8]
    if col_count == 6:
        return [1.4, 4.2, 2.6, 2.6, 2.4, 2.2]
    base = round(15.4 / col_count, 2)
    return [base] * col_count


def convert(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_styles(doc)
    add_page_number_footer(doc)

    title = lines[0].lstrip("# ").strip()
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(title)
    set_run_font(run, cjk="黑体", size=Pt(20), bold=True, color=RGBColor(0x1F, 0x4D, 0x78))

    idx = 1
    # 顶部元信息（引用块）
    meta: list[str] = []
    while idx < len(lines) and (lines[idx].strip().startswith(">") or lines[idx].strip() == ""):
        if lines[idx].strip().startswith(">"):
            meta.append(lines[idx].strip().lstrip("> ").strip())
        idx += 1
    for item in meta:
        body_paragraph(doc, item, indent=False, size=Pt(9.5))

    doc.add_paragraph()
    h1(doc, "目录")
    add_toc_field(doc)
    doc.add_page_break()

    in_code = False
    code_lines: list[str] = []

    while idx < len(lines):
        raw = lines[idx]
        line = raw.strip()

        img = re.match(r"^!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]+)\)$", line)
        if img:
            src = (md_path.parent / img.group("src")).resolve()
            alt = img.group("alt").strip()
            if src.is_file():
                para = doc.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.first_line_indent = Cm(0)
                para.add_run().add_picture(str(src), width=Cm(15.0))
                caption(doc, alt)
            else:
                body_paragraph(doc, f"（图片缺失：{img.group('src')}）", indent=False)
            idx += 1
            continue

        if line.startswith("```"):
            if in_code:
                code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            idx += 1
            continue
        if in_code:
            code_lines.append(raw)
            idx += 1
            continue

        if not line or line == "---":
            idx += 1
            continue

        if line.startswith("|"):
            rows, idx = parse_table_rows(lines, idx)
            header = split_row(rows[0])
            body = [split_row(r) for r in rows[2:]] if len(rows) > 1 else []
            add_table(doc, header, body, table_widths(len(header)))
            continue

        if line.startswith("#### "):
            h3(doc, line[5:].strip())
            idx += 1
            continue
        if line.startswith("### "):
            h2(doc, line[4:].strip())
            idx += 1
            continue
        if line.startswith("## "):
            h1(doc, line[3:].strip())
            idx += 1
            continue
        if line.startswith("# "):
            idx += 1
            continue

        if PLACEHOLDER_HINT in line and "图" in line:
            image_placeholder(doc, line)
            idx += 1
            continue

        if line.startswith("> "):
            note = line.lstrip("> ").strip()
            if note:
                note_block(doc, note)
            idx += 1
            continue

        if line.startswith("- "):
            body_paragraph(doc, "• " + line[2:].strip(), indent=False)
            idx += 1
            continue

        if re.match(r"^\d+\.\s+", line):
            body_paragraph(doc, line, indent=False)
            idx += 1
            continue

        body_paragraph(doc, line)
        idx += 1

    doc.save(docx_path)
    print(f"OK  {docx_path.relative_to(ROOT)}")


def main() -> None:
    for md_path, docx_path in TASKS:
        if not md_path.is_file():
            print(f"MISS {md_path}")
            continue
        convert(md_path, docx_path)


if __name__ == "__main__":
    main()
