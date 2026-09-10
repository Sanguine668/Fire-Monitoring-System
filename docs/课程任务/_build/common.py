"""课程文档生成公共库：A4 中文报告样式 + 表格/编号/图片/页脚辅助。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_CJK = "宋体"
HEAD_CJK = "黑体"
LATIN = "Times New Roman"
BODY_SIZE = Pt(10.5)
SMALL_SIZE = Pt(9)
ACCENT = RGBColor(0x1F, 0x4D, 0x78)
HEADER_FILL = "1F4D78"
ALT_FILL = "F2F5F9"


def set_run_font(run, cjk: str = BODY_CJK, latin: str = LATIN, size=None, bold=None, color=None, italic=None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cjk)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _style_font(style, cjk: str, latin: str, size: Pt, bold: bool, color: RGBColor | None = None):
    style.font.name = latin
    style.font.size = size
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    _style_font(normal, BODY_CJK, LATIN, BODY_SIZE, False)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(3)
    pf.space_before = Pt(0)

    heading_specs = {
        "Heading 1": (Pt(16), HEAD_CJK, True, Pt(14), Pt(8)),
        "Heading 2": (Pt(14), HEAD_CJK, True, Pt(11), Pt(6)),
        "Heading 3": (Pt(12), HEAD_CJK, True, Pt(8), Pt(4)),
    }
    for name, (size, cjk, bold, before, after) in heading_specs.items():
        style = doc.styles[name]
        _style_font(style, cjk, LATIN, size, bold, ACCENT)
        style.paragraph_format.space_before = before
        style.paragraph_format.space_after = after
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)

    for list_name in ("List Bullet", "List Number"):
        style = doc.styles[list_name]
        _style_font(style, BODY_CJK, LATIN, BODY_SIZE, False)
        style.paragraph_format.line_spacing = 1.4
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.left_indent = Cm(0.9)
        style.paragraph_format.first_line_indent = Cm(0)

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.6)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.5)


def add_page_number_footer(doc: Document) -> None:
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("第 ")
        set_run_font(run, size=SMALL_SIZE)
        field_begin = OxmlElement("w:fldChar")
        field_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        field_end = OxmlElement("w:fldChar")
        field_end.set(qn("w:fldCharType"), "end")
        run2 = p.add_run()
        run2._r.append(field_begin)
        run2._r.append(instr)
        run2._r.append(field_end)
        run3 = p.add_run(" 页")
        set_run_font(run3, size=SMALL_SIZE)


def add_toc_field(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "（在 Word 中按 F9 可更新目录）"
    run.append(text)
    fld.append(run)
    p._p.append(fld)


def new_document(title: str, subtitle: str, meta_rows: list[tuple[str, str]], doc_number: str = "") -> Document:
    doc = Document()
    configure_styles(doc)
    add_page_number_footer(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(90)
    r = p.add_run(title)
    set_run_font(r, cjk=HEAD_CJK, size=Pt(24), bold=True, color=ACCENT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    set_run_font(r, cjk=HEAD_CJK, size=Pt(14), bold=False, color=RGBColor(0x40, 0x40, 0x40))

    if doc_number:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"文档编号：{doc_number}")
        set_run_font(r, size=Pt(10.5))

    doc.add_paragraph()
    add_table(doc, ["项目", "内容"], [list(row) for row in meta_rows], [4.0, 10.0], center_body=False)
    doc.add_page_break()
    return doc


def h1(doc: Document, text: str) -> None:
    doc.add_heading(text, level=1)


def h2(doc: Document, text: str) -> None:
    doc.add_heading(text, level=2)


def h3(doc: Document, text: str) -> None:
    doc.add_heading(text, level=3)


def p(doc: Document, text: str, indent: bool = True, bold: bool = False, size=None, align=None):
    para = doc.add_paragraph()
    run = para.add_run(text)
    set_run_font(run, size=size or BODY_SIZE, bold=bold)
    if indent:
        para.paragraph_format.first_line_indent = Cm(0.74)
    if align is not None:
        para.alignment = align
    return para


def lead(doc: Document, text: str):
    return p(doc, text, indent=False, bold=True)


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        run = para.add_run(item)
        set_run_font(run)


def numbers(doc: Document, items: list[str]) -> None:
    for item in items:
        para = doc.add_paragraph(style="List Number")
        run = para.add_run(item)
        set_run_font(run)


def caption(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    set_run_font(run, size=Pt(9), bold=False, color=RGBColor(0x40, 0x40, 0x40))


def add_image(doc: Document, image_path: str | Path, width_cm: float, caption_text: str) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(image_path), width=Cm(width_cm))
    caption(doc, caption_text)


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_margins(table, top=60, bottom=60, start=100, end=100) -> None:
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("left", start), ("bottom", bottom), ("right", end)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tbl_pr.append(mar)


def _set_table_widths(table, widths_cm: list[float]) -> None:
    total_dxa = int(sum(widths_cm) * 567)
    tbl_pr = table._tbl.tblPr
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        table._tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_cm:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 567)))
        grid.append(col)
    table._tbl.insert(list(table._tbl).index(tbl_pr) + 1, grid)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[idx])


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths_cm: list[float],
    center_body: bool = False,
    font_size: Pt | None = None,
    zebra: bool = True,
):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_widths(table, widths_cm)
    _set_cell_margins(table)

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.first_line_indent = Cm(0)
        run = para.add_run(header)
        set_run_font(run, cjk=HEAD_CJK, size=font_size or Pt(9.5), bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _shade(cell, HEADER_FILL)

    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.first_line_indent = Cm(0)
            para.paragraph_format.line_spacing = 1.25
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if (center_body or c_idx == 0) else WD_ALIGN_PARAGRAPH.LEFT
            run = para.add_run(str(value))
            set_run_font(run, size=font_size or Pt(9.5))
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if zebra and r_idx % 2 == 1:
                _shade(cell, ALT_FILL)
    return table


def page_break(doc: Document) -> None:
    doc.add_page_break()
